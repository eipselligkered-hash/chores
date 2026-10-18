"""Write the CLI's JSON payload (stdin) to an .xlsx snapshot: All Tasks (values), Done log, This Weekend."""
import sys, json, datetime as dt, openpyxl
from openpyxl.styles import Font
p = json.load(sys.stdin); out = sys.argv[1]
D = lambda s: dt.date.fromisoformat(s) if s else None
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "All Tasks"
cols = [("room","Room"),("name","Task"),("next_due","Next due"),("status","Status"),("last_effective","Last done (effective)"),("plan","Plan"),
        ("last_done","Last done (typed)"),("frequency","Frequency"),("priority","Priority"),("effort_min","Effort (min)"),("planned_date","Planned date"),
        ("stagger_date","Stagger date (one-off)"),("due_weekday","Due weekday (override)"),("season","Season"),("notes","Notes")]
ws.append([h for _, h in cols])
for t in p["tasks"]:
    ws.append([D(t.get(k)) if k in ("next_due","last_effective","last_done","planned_date","stagger_date") else t.get(k) for k, _ in cols])
for r in ws.iter_rows(min_row=2):
    for c in r:
        if isinstance(c.value, dt.date): c.number_format = "ddd dd/mm/yyyy"
dl = wb.create_sheet("Done log"); dl.append(["Date","Task (Room: Task)","Note","Source"])
for c in sorted(p["completions"], key=lambda c: c["done_on"]):
    dl.append([D(c["done_on"]), f'{c.get("room")}: {c.get("task")}', c.get("note"), c.get("source")]); dl.cell(dl.max_row, 1).number_format = "dd/mm/yyyy"
tw = wb.create_sheet("This Weekend"); tw.append(["No.","Room","Task","Priority","Next due","Status","Plan","Effort (min)","Running total (capped)","Weekend"])
for i, r in enumerate(p["weekend"], 1):
    tw.append([i, r["room"], r["name"], r["priority"], D(r["next_due"]), r["status"], r["plan"], r["effort_min"], r["running"], r["bucket"]]); tw.cell(tw.max_row, 5).number_format = "ddd dd/mm"
for s in wb:
    for c in s[1]: c.font = Font(bold=True)
    s.freeze_panes = "A2"
wb.save(out); print(f"exported {len(p['tasks'])} tasks, {len(p['completions'])} completions, {len(p['weekend'])} weekend rows -> {out} (as of {p['today']})")
